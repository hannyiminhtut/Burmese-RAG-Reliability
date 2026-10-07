"""
build_docx_from_pdf.py
Converts docs/26042.pdf (compiled from paper_draft_v1.1_revised.tex) into a clean,
IEEE-compliant two-column editable .docx document: docs/26042.docx
Includes:
- Two-column layout with standard IEEE margins matching ICAIT2026 template
- Native OMML mathematical equations (Hit@k, MRR, Accuracy, P_A/R_A/F_1, Cochran's Q)
- High-resolution extracted figures (Fig. 1, Fig. 2, Fig. 3) properly styled and captioned
- Properly formatted IEEE tables (Table I, Table II, Table III, Table IV) with borders
- Clean references, headings, abstract, keywords, and author headers.
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

DOC_TEMPLATE = "docs/ICAIT2026-Template.docx"
OUT_DOCX = "docs/26042.docx"

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def style_ieee_table(table, col_widths, alignments):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        '<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        '<w:left w:val="none"/>'
        '<w:right w:val="none"/>'
        '<w:insideH w:val="none"/>'
        '<w:insideV w:val="none"/>'
        '</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    # Header bottom border
    if len(table.rows) > 0:
        for cell in table.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                '<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                '</w:tcBorders>'
            )
            tcPr.append(tcBorders)

    for row_idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml('<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
        if row_idx == 0:
            trPr.append(parse_xml('<w:tblHeader xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
        for col_idx, cell in enumerate(row.cells):
            cell.width = Inches(col_widths[col_idx])
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                align = alignments[col_idx]
                if align == 'C':
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif align == 'R':
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

def add_callout(p, text, bold_prefix=""):
    run = p.add_run(bold_prefix)
    run.bold = True
    p.add_run(text)

def create_omml_equation(eq_type):
    """Generate OMML XML for key equations"""
    if eq_type == "hitk":
        # Hit@k = (1 / N_A) \sum I(G_i \cap R_i^(k) \neq \emptyset)
        return (
            '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            '<m:oMath>'
            '<m:r><m:rPr><m:sty m:val="p"/></m:rPr><m:t>Hit@</m:t></m:r>'
            '<m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>k</m:t></m:r>'
            '<m:r><m:t> = </m:t></m:r>'
            '<m:f>'
            '<m:num><m:r><m:t>1</m:t></m:r></m:num>'
            '<m:den><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>N</m:t></m:r><m:r><m:rPr><m:scr m:val="script"/></m:rPr><m:t>A</m:t></m:r></m:den>'
            '</m:f>'
            '<m:nary>'
            '<m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/><m:grow m:val="1"/></m:naryPr>'
            '<m:sub><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>i</m:t></m:r><m:r><m:t>=1</m:t></m:r></m:sub>'
            '<m:sup><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>N</m:t></m:r><m:r><m:rPr><m:scr m:val="script"/></m:rPr><m:t>A</m:t></m:r></m:sup>'
            '<m:e>'
            '<m:r><m:rPr><m:sty m:val="p"/><m:b/></m:rPr><m:t>I</m:t></m:r>'
            '<m:d>'
            '<m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>'
            '<m:e>'
            '<m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>G</m:t></m:r>'
            '<m:r><m:rPr><m:scr m:val="script"/></m:rPr><m:t>i</m:t></m:r>'
            '<m:r><m:t> ∩ </m:t></m:r>'
            '<m:sSubSup>'
            '<m:e><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>R</m:t></m:r></m:e>'
            '<m:sub><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>i</m:t></m:r></m:sub>'
            '<m:sup><m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>k</m:t></m:r></m:e></m:d></m:sup>'
            '</m:sSubSup>'
            '<m:r><m:t> ≠ ∅</m:t></m:r>'
            '</m:e>'
            '</m:d>'
            '</m:e>'
            '</m:nary>'
            '<m:r><m:t>,       (1)</m:t></m:r>'
            '</m:oMath>'
            '</m:oMathPara>'
        )
    elif eq_type == "mrr":
        # MRR = (1 / N_A) \sum (1 / r_i)
        return (
            '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            '<m:oMath>'
            '<m:r><m:rPr><m:sty m:val="p"/></m:rPr><m:t>MRR</m:t></m:r>'
            '<m:r><m:t> = </m:t></m:r>'
            '<m:f>'
            '<m:num><m:r><m:t>1</m:t></m:r></m:num>'
            '<m:den><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>N</m:t></m:r><m:r><m:rPr><m:scr m:val="script"/></m:rPr><m:t>A</m:t></m:r></m:den>'
            '</m:f>'
            '<m:nary>'
            '<m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/><m:grow m:val="1"/></m:naryPr>'
            '<m:sub><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>i</m:t></m:r><m:r><m:t>=1</m:t></m:r></m:sub>'
            '<m:sup><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>N</m:t></m:r><m:r><m:rPr><m:scr m:val="script"/></m:rPr><m:t>A</m:t></m:r></m:sup>'
            '<m:e>'
            '<m:f>'
            '<m:num><m:r><m:t>1</m:t></m:r></m:num>'
            '<m:den><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>r</m:t></m:r><m:r><m:rPr><m:scr m:val="script"/></m:rPr><m:t>i</m:t></m:r></m:den>'
            '</m:f>'
            '</m:e>'
            '</m:nary>'
            '<m:r><m:t>,       (2)</m:t></m:r>'
            '</m:oMath>'
            '</m:oMathPara>'
        )
    elif eq_type == "acc":
        # Accuracy = (TP + TN) / (TP + TN + FP + FN)
        return (
            '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            '<m:oMath>'
            '<m:r><m:rPr><m:sty m:val="p"/></m:rPr><m:t>Accuracy</m:t></m:r>'
            '<m:r><m:t> = </m:t></m:r>'
            '<m:f>'
            '<m:num><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>TP</m:t></m:r><m:r><m:t> + </m:t></m:r><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>TN</m:t></m:r></m:num>'
            '<m:den><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>TP</m:t></m:r><m:r><m:t> + </m:t></m:r><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>TN</m:t></m:r><m:r><m:t> + </m:t></m:r><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>FP</m:t></m:r><m:r><m:t> + </m:t></m:r><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>FN</m:t></m:r></m:den>'
            '</m:f>'
            '<m:r><m:t>,       (3)</m:t></m:r>'
            '</m:oMath>'
            '</m:oMathPara>'
        )
    elif eq_type == "decision_metrics":
        # P_A = TP/(TP+FP),  R_A = TP/(TP+FN),  F_1 = 2 P_A R_A / (P_A + R_A)
        return (
            '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            '<m:oMath>'
            '<m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>P</m:t></m:r><m:r><m:rPr><m:scr m:val="script"/></m:rPr><m:t>A</m:t></m:r>'
            '<m:r><m:t> = </m:t></m:r>'
            '<m:f>'
            '<m:num><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>TP</m:t></m:r></m:num>'
            '<m:den><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>TP</m:t></m:r><m:r><m:t> + </m:t></m:r><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>FP</m:t></m:r></m:den>'
            '</m:f>'
            '<m:r><m:t>,   </m:t></m:r>'
            '<m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>R</m:t></m:r><m:r><m:rPr><m:scr m:val="script"/></m:rPr><m:t>A</m:t></m:r>'
            '<m:r><m:t> = </m:t></m:r>'
            '<m:f>'
            '<m:num><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>TP</m:t></m:r></m:num>'
            '<m:den><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>TP</m:t></m:r><m:r><m:t> + </m:t></m:r><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>FN</m:t></m:r></m:den>'
            '</m:f>'
            '<m:r><m:t>,   </m:t></m:r>'
            '<m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>F</m:t></m:r><m:r><m:rPr><m:scr m:val="script"/></m:rPr><m:t>1</m:t></m:r>'
            '<m:r><m:t> = </m:t></m:r>'
            '<m:f>'
            '<m:num><m:r><m:t>2</m:t></m:r><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>P</m:t></m:r><m:r><m:rPr><m:scr m:val="script"/></m:rPr><m:t>A</m:t></m:r><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>R</m:t></m:r><m:r><m:rPr><m:scr m:val="script"/></m:rPr><m:t>A</m:t></m:r></m:num>'
            '<m:den><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>P</m:t></m:r><m:r><m:rPr><m:scr m:val="script"/></m:rPr><m:t>A</m:t></m:r><m:r><m:t> + </m:t></m:r><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>R</m:t></m:r><m:r><m:rPr><m:scr m:val="script"/></m:rPr><m:t>A</m:t></m:r></m:den>'
            '</m:f>'
            '<m:r><m:t>,       (4)</m:t></m:r>'
            '</m:oMath>'
            '</m:oMathPara>'
        )
    elif eq_type == "cochran":
        # Q = \frac{(k-1) [ k \sum C_j^2 - T^2 ]}{ k T - \sum R_i^2 }
        return (
            '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            '<m:oMath>'
            '<m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>Q</m:t></m:r>'
            '<m:r><m:t> = </m:t></m:r>'
            '<m:f>'
            '<m:num>'
            '<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>k</m:t></m:r><m:r><m:t> - 1</m:t></m:r></m:e></m:d>'
            '<m:d><m:dPr><m:begChr m:val="["/><m:endChr m:val="]"/></m:dPr>'
            '<m:e>'
            '<m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>k</m:t></m:r>'
            '<m:nary>'
            '<m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/><m:grow m:val="1"/></m:naryPr>'
            '<m:sub><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>j</m:t></m:r><m:r><m:t>=1</m:t></m:r></m:sub>'
            '<m:sup><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>k</m:t></m:r></m:sup>'
            '<m:e>'
            '<m:sSubSup>'
            '<m:e><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>C</m:t></m:r></m:e>'
            '<m:sub><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>j</m:t></m:r></m:sub>'
            '<m:sup><m:r><m:t>2</m:t></m:r></m:sup>'
            '</m:sSubSup>'
            '</m:e>'
            '</m:nary>'
            '<m:r><m:t> - </m:t></m:r>'
            '<m:sSup><m:e><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>T</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>'
            '</m:e>'
            '</m:d>'
            '</m:num>'
            '<m:den>'
            '<m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>k</m:t></m:r><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>T</m:t></m:r>'
            '<m:r><m:t> - </m:t></m:r>'
            '<m:nary>'
            '<m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/><m:grow m:val="1"/></m:naryPr>'
            '<m:sub><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>i</m:t></m:r><m:r><m:t>=1</m:t></m:r></m:sub>'
            '<m:sup><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>N</m:t></m:r></m:sup>'
            '<m:e>'
            '<m:sSubSup>'
            '<m:e><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>R</m:t></m:r></m:e>'
            '<m:sub><m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>i</m:t></m:r></m:sub>'
            '<m:sup><m:r><m:t>2</m:t></m:r></m:sup>'
            '</m:sSubSup>'
            '</m:e>'
            '</m:nary>'
            '</m:den>'
            '</m:f>'
            '<m:r><m:t>,       (5)</m:t></m:r>'
            '</m:oMath>'
            '</m:oMathPara>'
        )

print("Helper functions defined successfully")
