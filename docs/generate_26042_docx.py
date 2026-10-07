"""
generate_26042_docx.py
Converts 26042.pdf / paper_draft_v1.1_revised.tex into the official editable Word document:
  docs/26042.docx
Key properties:
- Uses the official IEEE conference template (docs/ICAIT2026-Template.docx)
- Page size A4 (8.27 x 11.69 in), top/bottom/left/right standard margins
- Authors header: centered 2-column table with clean styling
- Body text: 2-column format (0.25 in spacing between columns)
- Headings: Roman numerals (I. Introduction, II. Related Work, etc.)
- Native OMML Equation objects for equations (1)-(5)
- All 4 IEEE tables with borders, correct widths, numbers, and clean cell padding
- All 3 high-res figures (Fig. 1, Fig. 2, Fig. 3) with IEEE captions
- Complete references formatted with IEEE reference style
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
import docs.docx_helpers as dh

def build_docx():
    template_path = "docs/ICAIT2026-Template.docx"
    output_path = "docs/26042.docx"
    doc = docx.Document(template_path)
    b = doc._body._element

    # 1. Clear body children except body sectPr
    body_sectPr = b.xpath('./w:sectPr')[0]
    for child in list(b):
        if child is not body_sectPr:
            b.remove(child)

    # Configure body_sectPr to 2 columns
    body_cols = body_sectPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cols')
    if body_cols is not None:
        body_cols.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num', '2')
        body_cols.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}space', '360')
    else:
        body_sectPr.append(parse_xml('<w:cols xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:num="2" w:space="360"/>'))

    # Helper function for adding paragraphs
    def add_p(text="", style="Body Text", space_before=0, space_after=4, align=None):
        p = doc.add_paragraph(style=style)
        if text:
            p.text = text
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.05
        if align is not None:
            p.alignment = align
        return p

    def add_h1(title):
        p = add_p(title, style="Heading 1", space_before=8, space_after=3)
        return p

    def add_h2(title):
        p = add_p(title, style="Heading 2", space_before=5, space_after=2)
        return p

    def add_fig(img_path, caption, width_in=3.35):
        if os.path.exists(img_path):
            p_img = add_p(style="Normal", space_before=6, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
            p_img.add_run().add_picture(img_path, width=Inches(width_in))
        p_cap = add_p(style="figure caption", space_before=2, space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        # Parse Fig. X. prefix
        m = re.match(r'^(Fig\.\s*\d+\.)\s*(.*)$', caption)
        if m:
            r1 = p_cap.add_run(m.group(1) + " ")
            r1.font.size = Pt(8)
            r2 = p_cap.add_run(m.group(2))
            r2.font.size = Pt(8)
        else:
            r = p_cap.add_run(caption)
            r.font.size = Pt(8)

    def add_eq(eq_type):
        p = add_p(style="Normal", space_before=3, space_after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
        omml_xml = dh.create_omml_equation(eq_type)
        p._p.append(parse_xml(omml_xml))

    # ==========================
    # TITLE (1-column section)
    # ==========================
    p_title = add_p("Language-Conditioned Abstention in Burmese-English RAG: Retrieval–Decision Divergence Across Query Languages",
                    style="paper title", space_before=0, space_after=12, align=WD_ALIGN_PARAGRAPH.CENTER)

    # AUTHORS TABLE (2 authors)
    t_auth = doc.add_table(rows=1, cols=2)
    t_auth.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_auth.autofit = False
    tblPr = t_auth._tbl.tblPr
    tblPr.append(parse_xml(
        '<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/>'
        '<w:insideH w:val="none"/><w:insideV w:val="none"/>'
        '</w:tblBorders>'
    ))
    
    # Author 1
    c0 = t_auth.rows[0].cells[0]
    c0.width = Inches(3.4)
    p0 = c0.paragraphs[0]
    p0.style = "Author"
    p0.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_after = Pt(2)
    r_a1 = p0.add_run("Han Nyi Min Htut\n")
    r_a1.bold = True
    p0.add_run("University of Information Technology\nYangon, Myanmar\n")
    p0.add_run("hannyiminhtut@uit.edu.mm")

    # Author 2
    c1 = t_auth.rows[0].cells[1]
    c1.width = Inches(3.4)
    p1 = c1.paragraphs[0]
    p1.style = "Author"
    p1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(2)
    r_a2 = p1.add_run("Thet Thet Zin\n")
    r_a2.bold = True
    p1.add_run("University of Information Technology\nYangon, Myanmar\n")
    p1.add_run("thetthetzin@uit.edu.mm")

    # Continuous section break to switch to 2 columns
    p_break = doc.add_paragraph()
    p_break.paragraph_format.space_before = Pt(0)
    p_break.paragraph_format.space_after = Pt(0)
    p_break._p.get_or_add_pPr().append(parse_xml(
        '<w:sectPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:type w:val="continuous"/>'
        '<w:cols w:space="720"/>'
        '</w:sectPr>'
    ))

    # ==========================
    # ABSTRACT & KEYWORDS (2-column section)
    # ==========================
    p_abs = add_p(style="Abstract", space_before=4, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    r_ab_label = p_abs.add_run("Abstract—")
    r_ab_label.bold = True
    p_abs.add_run(
        "Retrieval-augmented generation (RAG) systems for university regulations must answer supported questions and abstain "
        "when evidence is unavailable. This paper reports a controlled structural and decision-level evaluation of one local "
        "RAG configuration over a frozen University of Information Technology academic-regulation document, using English (EN), "
        "Burmese (MY), and Burmese-English code-mixed (MIX) queries. The benchmark has 79 paired semantic intents and 237 variants, "
        "of which 222 are held out. Retrieval uses intfloat/multilingual-e5-small with FAISS inner-product search, and the top five "
        "chunks are supplied to gemma3:4b-it-q4_K_M. On 180 answerable held-out variants, overall Hit@5 was 87.22%. Accuracy of the "
        "structured ANSWER/ABSTAIN decision was 73.42%. This value measures whether the system chose to answer, not whether its answer "
        "was correct. MY decision accuracy was 60.81%, compared with 79.73% for both EN and MIX. The main language-conditioned issue was "
        "false abstention: 45.00% for MY versus 20.00% for EN and 16.67% for MIX. Retrieval Hit@k differences were not statistically "
        "significant, whereas the decision difference was (p = 0.0050). Exploratory diagnostics found 13 MY false abstentions with gold-page "
        "evidence ranked first. False-answer differences, based on only 14 unanswerable intents per language, were not significant and are "
        "underpowered. In total, 95.50% of records satisfied the structural output contract; this checks format and citation binding, not "
        "factual correctness. Semantic answer correctness, groundedness, and hallucination were not evaluated, so the findings apply to "
        "the evaluated corpus and configuration."
    )

    p_kw = add_p(style="Keywords", space_before=2, space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    r_kw_label = p_kw.add_run("Index Terms—")
    r_kw_label.bold = True
    p_kw.add_run("retrieval-augmented generation, Burmese, code mixing, abstention, university regulations, multilingual retrieval, reliability")

    # ==========================
    # I. INTRODUCTION
    # ==========================
    add_h1("I. Introduction")
    add_p(
        "Academic-credit rules affect registration, examinations, progression, and graduation. Students often need to locate "
        "a specific rule rather than read an entire regulation document. A RAG interface can make that search easier, provided "
        "that its answers remain tied to the institutional source. Retrieval alone is not enough: the system must also judge "
        "whether the retrieved text supports an answer, decline unsupported questions, respond in an appropriate language, and "
        "identify the evidence used."
    )
    add_p(
        "The UIT source is written mainly in Burmese but regularly uses English academic terms. Student questions may therefore "
        "be written in English, in Burmese, or in a mixture of both. Even when the requested fact is unchanged, these surface "
        "forms can lead to different embedding matches and different generator decisions. An overall score would hide such "
        "cases, so the three language forms are compared within each underlying information need."
    )
    add_p("This study examines a frozen local RAG configuration at three separately reported measurement levels (Fig. 1):")
    
    # Bullet points for 3 levels
    for b_item, b_desc in [
        ("retrieval: ", "whether gold-page evidence is returned;"),
        ("decision: ", "whether the generator chooses ANSWER or ABSTAIN in agreement with the benchmark answerability label;"),
        ("structural contract: ", "whether the output satisfies the required format and chunk–page citation binding.")
    ]:
        p_bp = add_p(style="bullet list", space_before=1, space_after=1)
        r_bp = p_bp.add_run("•  " + b_item)
        r_bp.italic = True
        p_bp.add_run(b_desc)

    add_p(
        "None of these levels measures whether a generated answer is semantically correct or grounded; that dimension is outside "
        "the scope of the reported evaluation. The study asks:"
    )

    for q_num, q_text in [
        ("1) ", "How effectively does the frozen multilingual retriever return gold-page evidence for answerable EN, MY, and MIX questions?"),
        ("2) ", "How accurately does the generator choose ANSWER or ABSTAIN across the three language conditions?"),
        ("3) ", "Which reproducible structural failure patterns account for observed language-condition differences?")
    ]:
        p_q = add_p(style="Body Text", space_before=1, space_after=1)
        p_q.paragraph_format.left_indent = Inches(0.15)
        p_q.add_run(q_num).bold = True
        p_q.add_run(q_text)

    add_p(
        "The resulting contribution is a frozen university-domain benchmark, a locally reproducible retrieval and generation pipeline, "
        "and a paired analysis of EN, MY, and MIX variants. The implementation also records the chunk–page relationship used by each "
        "citation, saves every completed record, and retains validation and infrastructure failures in the analysis. The work is "
        "presented as a focused, configuration-specific reliability study rather than as a general claim about Burmese RAG systems."
    )

    # FIGURE 1
    add_fig("docs/fig1_pipeline.png", "Fig. 1. Evaluated pipeline. Retrieval, structured decision, and structural contract are reported as separate measurement levels; semantic answer quality is not measured.", width_in=3.4)

    # ==========================
    # II. RELATED WORK
    # ==========================
    add_h1("II. Related Work")
    add_p(
        "RAG couples a parametric generator with evidence retrieved from an external collection [1]. Because a single end-to-end "
        "score can conceal where an error occurred, prior evaluation work separates retrieval quality, use of context, and generated "
        "output [3]. We follow this separation."
    )
    add_p(
        "The retriever uses multilingual E5 [2]. Language condition can influence both stages of a multilingual RAG pipeline. "
        "Code mixing may help cross-lingual alignment in one case and weaken direct query–passage matching in another [4]. "
        "Retriever and generator language preferences can also diverge, so better retrieval does not guarantee better generation [5]. "
        "This work motivates the paired EN/MY/MIX comparison but does not predict which condition should perform best for Burmese."
    )
    add_p(
        "Abstention is increasingly treated as a measurable LLM behavior [6], and selective abstention has been proposed for insufficient "
        "or misaligned knowledge [7]. We do not train an abstention mechanism; we compare the frozen generator's decision with the "
        "benchmark label. Detecting hallucination requires semantic judgement [8], so citation validity is used here only as a "
        "structural measure."
    )
    add_p(
        "Burmese has limited coverage in NLP evaluation, although newer benchmarks examine it more systematically [9]. Low-resource "
        "cross-lingual retrieval augmentation [10], Bangla RAG [11], and university-domain QA [12] all suggest that English results "
        "should not be assumed to transfer unchanged to Burmese."
    )

    # ==========================
    # III. DATASET AND SCOPE
    # ==========================
    add_h1("III. Dataset and Scope")
    add_h2("A. Source Corpus")
    add_p(
        "The source is UIT's 24-page Undergraduate Academic Regulations (4 Years Program), dated June 2025 and archived as UIT-SRC-025 [13]. "
        "It is written in Burmese and retains English academic terminology. Page 1 is the title page, and pages 2–3 are the table of contents. "
        "These pages were excluded from the index because they contain no operative rules and repeat section wording that could create duplicate "
        "lexical matches. Corpus version UIT-AC-v1.1 contains 148 retrieval chunks from pages 4–24. The text, chunk IDs, page metadata, "
        "and hashes were frozen before held-out generation. This corpus covers one official document, not the full UIT website, so the results "
        "concern the indexed academic-credit rules only."
    )

    add_h2("B. Benchmark")
    add_p(
        "The benchmark contains 79 semantic intents, each represented by EN, MY, and MIX variants, for 237 questions in total (189 answerable, "
        "48 unanswerable). For an answerable intent, the record includes a gold answer, source ID, page, and regulation section. An unanswerable "
        "intent requests information that is absent from the frozen document; it has no gold-evidence page and expects ABSTAIN. The three "
        "variants of an intent keep the requested fact and answerability label fixed. Automated checks confirmed unique variant IDs, one complete "
        "language triplet per intent, consistent labels within each triplet, and complete evidence fields for answerable records."
    )
    add_p(
        "Five intents (Q002, Q008, Q024, Q029, Q057) were frozen as development data. The remaining 74 intents form the held-out test set "
        "(Table I). No development record was used in held-out retrieval, generation, statistics, or error analysis, and held-out observations "
        "were not used for tuning."
    )

    # TABLE I
    p_t1_head = add_p(style="table head", space_before=6, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    p_t1_head.add_run("TABLE I\nDataset and Split Summary")
    t1 = doc.add_table(rows=4, cols=6)
    t1_headers = ["Split", "Intents", "Variants", "Ans.", "Unans.", "EN/MY/MIX"]
    t1_data = [
        ["Development", "5", "15", "9", "6", "5/5/5"],
        ["Held-out", "74", "222", "180", "42", "74/74/74"],
        ["Total", "79", "237", "189", "48", "79/79/79"]
    ]
    for c_i, h in enumerate(t1_headers):
        t1.rows[0].cells[c_i].paragraphs[0].text = h
        t1.rows[0].cells[c_i].paragraphs[0].runs[0].bold = True
    for r_i, row in enumerate(t1_data):
        for c_i, val in enumerate(row):
            t1.rows[r_i+1].cells[c_i].paragraphs[0].text = val
    dh.style_ieee_table(t1, [0.8, 0.45, 0.5, 0.45, 0.5, 0.7], ['L', 'R', 'R', 'R', 'R', 'C'])

    # ==========================
    # IV. METHODOLOGY
    # ==========================
    add_h1("IV. Methodology")
    add_h2("A. Retrieval")
    add_p(
        "Chunk text is embedded with intfloat/multilingual-e5-small using the E5 passage prefix passage: , and queries use the prefix query: . "
        "Vectors are normalized, and cosine similarity is implemented through FAISS IndexFlatIP. The retriever returns exactly five chunks "
        "per question. A retrieved chunk is counted as relevant when any page in its metadata page list intersects the gold-evidence pages. "
        "Hit@k and mean reciprocal rank (MRR) are calculated only for answerable questions."
    )

    add_h2("B. Generation")
    add_p(
        "Generation uses the local Ollama model gemma3:4b-it-q4_K_M with prompt version uit-context-only-json-v3 and parser version "
        "uit-citation-object-parser-v3. The CPU-safe profile uses num_ctx=4096, num_predict=512, temperature 0, top_p=0.9, top_k=40, "
        "seed 42, and sequential requests."
    )
    add_p(
        "Each prompt contains the frozen question and the exact retrieved Top-5 chunks, each bound to its page list. The structured "
        "output contains a decision, an answer, and citations. ANSWER requires at least one citation that matches a supplied chunk–page "
        "pair. ABSTAIN requires no citation and must match a deterministic English abstention sentence. The contract is therefore "
        "language-asymmetric: ANSWER text must follow the query language (Burmese for MY), whereas ABSTAIN uses the same English sentence "
        "in every condition. Each completed record is checkpointed immediately. Transient infrastructure failures may receive one bounded "
        "retry; parsing, citation, and validation failures are not retried."
    )

    add_h2("C. Evaluation Boundaries")
    add_p(
        "Two terms are used narrowly throughout. Decision accuracy is the agreement between the generated decision field and the "
        "answerability label; a correct ANSWER decision can still contain a wrong or unsupported answer. Contract validity is a "
        "deterministic structural check of required fields, chunk–page citation binding, and abstention wording; it does not assess "
        "factual correctness, groundedness, or whether a citation supports a claim. Semantic answer assessment was not conducted "
        "(Section VIII)."
    )

    add_h2("D. Metrics and Statistical Analysis")
    add_p("For answerable question i, let Gi be its gold-page set and Ri^(k) the pages returned in the first k ranks:")

    # EQUATIONS (1) and (2)
    add_eq("hitk")
    add_eq("mrr")

    add_p(
        "where NA = 180, ri is the rank of the first relevant chunk, and 1/ri = 0 when no relevant chunk occurs in the Top-5. "
        "Treating ABSTAIN as the positive class, the decision measures are:"
    )

    # EQUATIONS (3) and (4)
    add_eq("acc")
    add_eq("decision_metrics")

    add_p(
        "where TP is a correct abstention, FP a false abstention, TN a correct answer decision, and FN a false answer. "
        "EN, MY, and MIX variants of an intent are paired. With Cj the successes in condition j, Ri the successes for intent i, "
        "T = ∑ Cj, and k = 3, Cochran's omnibus statistic is:"
    )

    # EQUATION (5)
    add_eq("cochran")

    add_p(
        "compared with χ²₂. For decision correctness, (C_EN, C_MY, C_MIX) = (59, 45, 59), T = 163, and ∑ R_i² = 415, giving Q = 10.595 "
        "and p = 0.0050. When Q is significant, two-sided exact McNemar tests with Holm adjustment compare language pairs, and 95% Wilson "
        "intervals quantify uncertainty. The retrieval and false-abstention populations contain 60 answerable intents per condition. "
        "The false-answer population contains only 14 unanswerable intents per condition. This gives low statistical power, so a "
        "non-significant false-answer result is not evidence that false-answer risk is equal across languages."
    )

    # ==========================
    # V. IMPLEMENTATION CHALLENGES
    # ==========================
    add_h1("V. Implementation Challenges and Design Innovations")
    add_h2("A. Multilingual Evidence Alignment")
    add_p(
        "The corpus mixes Burmese text with English terms such as credit, semester, and registration. During development, the EN, MY, "
        "and MIX versions of the same question sometimes produced different chunk rankings. The three variants were therefore grouped "
        "by intent, and each chunk received a stable ID and page list. A gold-page match was counted only as a retrieval hit, not as "
        "evidence that the answer was correct."
    )

    add_h2("B. Structured Output and Citation Control")
    add_p(
        "Early outputs sometimes contained invalid page values, citation objects copied from the example schema, or text appended to an "
        "ABSTAIN response. Each chunk was therefore supplied with its page list, and the model was required to return a fixed JSON object. "
        "A citation was accepted only when its page belonged to the cited chunk. Failed records preserved the original response and the "
        "stage at which validation failed."
    )

    add_h2("C. Resource-Constrained Local Execution")
    add_p(
        "CPU generation was slow, and the Ollama runner occasionally stopped because of memory limits. Execution was therefore restricted "
        "to one request and one loaded model at a time, with fixed context and output limits. Each record was saved before the next request, "
        "so interrupted runs could resume without overwriting failed records."
    )

    # ==========================
    # VI. RESULTS
    # ==========================
    add_h1("VI. Results")
    add_p("Retrieval (Table II), decision (Tables III–IV), and contract results are reported separately. They should not be combined or read as answer correctness.")

    add_h2("A. Retrieval Performance")
    # TABLE II
    p_t2_head = add_p(style="table head", space_before=6, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    p_t2_head.add_run("TABLE II\nRetrieval Results on Answerable Held-Out Variants")
    t2 = doc.add_table(rows=5, cols=5)
    t2_headers = ["Language", "Hit@1", "Hit@3", "Hit@5", "MRR"]
    t2_data = [
        ["Overall", "52.78%", "77.22%", "87.22%", "0.6577"],
        ["EN", "51.67%", "75.00%", "86.67%", "0.6444"],
        ["MY", "46.67%", "75.00%", "81.67%", "0.6039"],
        ["MIX", "60.00%", "81.67%", "93.33%", "0.7247"]
    ]
    for c_i, h in enumerate(t2_headers):
        t2.rows[0].cells[c_i].paragraphs[0].text = h
        t2.rows[0].cells[c_i].paragraphs[0].runs[0].bold = True
    for r_i, row in enumerate(t2_data):
        for c_i, val in enumerate(row):
            t2.rows[r_i+1].cells[c_i].paragraphs[0].text = val
    dh.style_ieee_table(t2, [0.85, 0.6, 0.6, 0.6, 0.75], ['L', 'R', 'R', 'R', 'R'])

    add_p(
        "MIX had the highest observed retrieval values and MY the lowest. However, none of the paired Hit@k comparisons was "
        "statistically significant; even the closest result, Hit@5, had p = 0.0990. Hit@k is a page-intersection measure and does "
        "not prove that the retrieved text was sufficient to answer."
    )

    add_h2("B. Structured Decision Performance")
    # TABLE III
    p_t3_head = add_p(style="table head", space_before=6, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    p_t3_head.add_run("TABLE III\nFrozen Label Versus Structured Decision")
    t3 = doc.add_table(rows=3, cols=3)
    t3_headers = ["Frozen label", "ANSWER", "ABSTAIN"]
    t3_data = [
        ["Answerable", "131", "49"],
        ["Unanswerable", "10", "32"]
    ]
    for c_i, h in enumerate(t3_headers):
        t3.rows[0].cells[c_i].paragraphs[0].text = h
        t3.rows[0].cells[c_i].paragraphs[0].runs[0].bold = True
    for r_i, row in enumerate(t3_data):
        for c_i, val in enumerate(row):
            t3.rows[r_i+1].cells[c_i].paragraphs[0].text = val
    dh.style_ieee_table(t3, [1.3, 1.0, 1.0], ['L', 'R', 'R'])

    add_p(
        "The generator produced 141 ANSWER and 81 ABSTAIN decisions (Table III). Decision accuracy was 163/222 (73.42%). This value "
        "measures only whether the decision field matched the answerability label; it is not answer-generation accuracy. Abstention "
        "precision was 32/81 (39.51%), recall was 32/42 (76.19%), and F1 was 52.03%. The system therefore detected most unanswerable "
        "variants but also rejected many answerable ones."
    )
    add_p(
        "Table IV gives decision results by language. Decision accuracy was 79.73% for EN and MIX but 60.81% for MY. Cochran's Q "
        "confirmed a language difference (Q(2) = 10.595, p = 0.0050), and Holm-adjusted pairwise tests isolated MY as significantly "
        "lower than both EN (p_adj = 0.0376) and MIX (Fig. 2)."
    )

    # TABLE IV
    p_t4_head = add_p(style="table head", space_before=6, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    p_t4_head.add_run("TABLE IV\nDecision Results by Language (Not Answer Correctness)")
    t4 = doc.add_table(rows=4, cols=4)
    t4_headers = ["Language", "Accuracy", "False abstain", "False answer"]
    t4_data = [
        ["EN", "79.73%", "20.00%", "3/14"],
        ["MY", "60.81%", "45.00%", "2/14"],
        ["MIX", "79.73%", "16.67%", "5/14"]
    ]
    for c_i, h in enumerate(t4_headers):
        t4.rows[0].cells[c_i].paragraphs[0].text = h
        t4.rows[0].cells[c_i].paragraphs[0].runs[0].bold = True
    for r_i, row in enumerate(t4_data):
        for c_i, val in enumerate(row):
            t4.rows[r_i+1].cells[c_i].paragraphs[0].text = val
    dh.style_ieee_table(t4, [0.85, 0.85, 0.95, 0.75], ['L', 'R', 'R', 'R'])
    
    p_t4_fn = add_p("Accuracy: 74 variants per language. False abstain: 60 answerable variants per language. False answer: 14 unanswerable variants per language.",
                    style="table footnote", space_before=1, space_after=4)
    p_t4_fn.runs[0].font.size = Pt(7.5)

    # FIGURE 2
    add_fig("docs/fig2_decision.png", "Fig. 2. Structured decision accuracy with 95% Wilson intervals. Accuracy concerns the ANSWER/ABSTAIN decision, not semantic answer correctness.", width_in=3.4)

    add_h2("C. False Abstention and False Answer")
    add_p(
        "Among the 60 answerable intents per condition, false-abstention rates were 20.00% for EN, 45.00% for MY, and 16.67% for MIX "
        "(Q(2) = 16.188, p = 0.00031). Holm-adjusted exact McNemar tests showed higher MY false abstention than EN (p_adj = 0.01185) "
        "and MIX (p_adj = 0.00146), whereas EN and MIX did not differ (p_adj = 0.7905) (Fig. 3)."
    )
    add_p(
        "Among 14 unanswerable intents per condition, false-answer counts were three for EN, two for MY, and five for MIX. The paired "
        "comparison was not significant (p = 0.2466). With only 14 intents per condition, differences of one to three intents cannot "
        "be reliably interpreted, and the result does not show equal false-answer risk. The observed ordering is therefore not "
        "interpreted further."
    )

    # FIGURE 3
    add_fig("docs/fig3_abstention.png", "Fig. 3. False-abstention rate with 95% Wilson intervals. False abstention means an answerable label paired with structured ABSTAIN.", width_in=3.4)

    add_h2("D. Contract and System Reliability")
    add_p(
        "Contract validity is a structural measure. It checks required fields, chunk–page citation pairs, and exact abstention wording, "
        "but not factual correctness or whether a citation supports the generated claim. Contract validation passed for 212 of 222 records "
        "(95.50%). Because the decision is parsed before the contract checks, a record can appear in the confusion matrix and still fail "
        "validation. Exact abstention wording was valid for 78 of 81 ABSTAIN outputs, and mechanical validation passed for 134 of 141 "
        "ANSWER outputs. Five records carried an infrastructure-event flag; all 222 records were retained. Mean generation time was "
        "69.06 s per record (about 4.26 h in total)."
    )

    add_h2("E. Structural Error Patterns")
    add_p(
        "The structural inventory contained 49 false abstentions, 10 false answers, 23 answerable Hit@5 misses, 10 contract-invalid records, "
        "and five records with an infrastructure-event flag. These overlapping flags affected 77 unique variants. Of the 49 false abstentions, "
        "34 coincided with a Top-5 gold-page hit and 15 with a miss. At the intent level, 37/74 intents had different EN/MY/MIX decisions. "
        "The most frequent pattern was EN=ANSWER, MY=ABSTAIN, MIX=ANSWER (17 intents). MY falsely abstained while both EN and MIX were "
        "correct for 16 intents."
    )

    add_h2("F. Qualitative Error Analysis")
    add_p("The following post-hoc observations are descriptive. They were computed deterministically from the frozen retrieval and generation records, and no semantic labels were assigned.")
    add_p(
        "Retrieval rank. Thirteen MY false abstentions occurred when a gold-page chunk was ranked first, compared with six for EN and "
        "four for MIX. Retrieval similarity also did not separate MY abstentions from MY answers (median top-1 score 0.860 versus 0.859)."
    )
    add_p(
        "Retrieved context. In the 16 intents where only MY falsely abstained, the MY Top-5 contained no gold page in seven cases; in six "
        "of these, both EN and MIX retrieved one. In the other nine, MY retrieved a gold page, at rank one in four cases. No MY Top-5 set "
        "was identical to its EN or MIX counterpart (median Jaccard overlap 0.29). For example, for “What is a Head of Department?” (Q056), "
        "all three variants retrieved a gold chunk at rank one; EN and MIX answered, whereas the formal Burmese variant returned the English "
        "abstention sentence. For the dismissal-condition question Q018, by contrast, only MY missed the gold page and only MY abstained."
    )
    add_p(
        "Prompt length and output language. All prompts were far below the 4,096-token context (median 1,009 EN, 1,085 MY, 1,122 MIX; "
        "maximum 1,795), so truncation is unlikely to explain the MY pattern. All 35 MY answers were written in Burmese, consistent with "
        "the language requirement."
    )
    add_p(
        "False answers. Five of the ten false answers stated in their text that the requested detail was absent (e.g., Q023-MIX, which asks "
        "for the LMS URL), although the decision field was ANSWER. This indicates an inconsistency between the decision field and the answer "
        "text rather than a confident fabrication. Other false answers, such as the program-subject question Q030 in all three languages, "
        "gave substantive content for information absent from the document."
    )
    add_p(
        "Output quality hidden by decision accuracy. Thirteen of the 55 MIX answers contained characters from unrelated scripts such as "
        "Telugu, Kannada, Greek, and CJK. The high MIX decision accuracy therefore does not by itself indicate clean or correct answers."
    )

    # ==========================
    # VII. DISCUSSION
    # ==========================
    add_h1("VII. Discussion")
    add_p(
        "The language pattern differed between retrieval and decision outcomes. MIX had the highest observed Hit@k values, but no paired "
        "retrieval test was significant, whereas MY decisions showed significantly more false abstention. The current experiments do not "
        "identify the cause. The evidence bears differently on four candidate explanations."
    )
    add_p(
        "Retrieval ranking. Ranking is unlikely to be the sole explanation, because retrieval differences were not significant and "
        "13 MY false abstentions occurred with a gold-page chunk ranked first."
    )
    add_p(
        "Retrieved-context completeness. A gold-page intersection does not guarantee that a chunk contains the complete rule. MY also "
        "retrieved a different Top-5 set from EN and MIX in every MY-only failure, and in seven of the 16 it retrieved no gold page. "
        "Context differences therefore remain a plausible contributor that only semantic evidence-sufficiency review could confirm."
    )
    add_p(
        "Generator language preference and contract asymmetry. The contract requires a Burmese answer for MY but a fixed English abstention "
        "sentence. For a small quantized generator, producing the English sentence may be easier than producing a Burmese answer, which could "
        "favour abstention. This hypothesis is consistent with reported differences between retriever and generator language preferences [5], "
        "but it was not tested here."
    )
    add_p(
        "Query formulation. MY variants use formal written Burmese, whereas MIX variants keep English key terms (e.g., CGPA, Re-exam) "
        "that also appear in the source. The surface form may thus affect both retrieval and the decision, and the present design cannot "
        "separate it from the other factors."
    )
    add_p(
        "Language property or configuration property? The evidence supports a configuration-specific finding for multilingual-e5-small, "
        "gemma3:4b-it-q4_K_M, and prompt uit-context-only-json-v3. It does not show that Burmese queries inherently cause abstention. "
        "A different generator, a Burmese abstention string, or a larger embedder could change the pattern. A context-swap experiment, "
        "giving MY questions the EN Top-5 context, would directly separate query-language from retrieved-context effects. The 34 gold-hit "
        "false abstentions provide a defined subset for later semantic review."
    )
    add_p(
        "Abstention recall was 76.19%, but precision was 39.51% because 49 answerable variants received ABSTAIN. For student use, false "
        "answers risk unsupported guidance, while false abstentions reduce coverage. No deployment threshold was selected, because the "
        "held-out set was not used for tuning. The 95.50% contract validity measures output structure, citation binding, and abstention "
        "wording only; a contract-valid record may still be incorrect, ungrounded, or script-contaminated (Section VI-F)."
    )
    add_p(
        "Possible transfer. Similar retrieval–decision divergence may appear wherever source documents mix a low-resource script with "
        "English terminology, the generator is weaker in the query language than in English, and abstention is cheap to produce. "
        "Institutional documents in other Southeast and South Asian languages are plausible candidates, but whether the pattern transfers "
        "remains an open empirical question."
    )

    # ==========================
    # VIII. LIMITATIONS
    # ==========================
    add_h1("VIII. Limitations")
    add_p(
        "First, qualified independent semantic reviewers were unavailable. The study therefore cannot estimate semantic answer "
        "correctness, groundedness, hallucination, numeric correctness, or claim-level citation support. The 73.42% decision accuracy "
        "and 95.50% contract validity are structural indicators, and the proportion of useful, correct answers may be lower. "
        "The contribution is accordingly framed as a structural and decision-level reliability study."
    )
    add_p(
        "Second, the corpus is a single 24-page UIT regulation document, so results may not generalize to other regulations, institutions, "
        "layouts, domains, or Burmese RAG systems in general. Third, only 14 unanswerable intents per language limit the power of false-answer "
        "comparisons; the non-significant result should not be read as equal risk. Fourth, one embedder, one generator, and one frozen prompt "
        "were evaluated, so the MY pattern may be configuration-specific. Fifth, the explanations in Section VII rest on descriptive, "
        "post-hoc associations, and no intervention such as a context swap was performed, so no causal attribution is made. Sixth, gold "
        "evidence is page based, which is less precise than claim-level annotation. Finally, runtimes reflect CPU-only local inference."
    )

    # ==========================
    # IX. ETHICAL AND REPRODUCIBILITY
    # ==========================
    add_h1("IX. Ethical and Reproducibility Considerations")
    add_p(
        "Only publicly accessible official UIT material was included. Personal, authentication-protected, unofficial, and unrelated "
        "information was excluded. The system should not be deployed as an authoritative substitute for official regulations, and users "
        "must be directed to the source when decisions have academic consequences."
    )
    add_p(
        "Corpus, benchmark, split, index, prompt, parser, resource profile, model digest, generation outputs, analysis protocols, "
        "result tables, and figures are versioned and hashed. Validation failures and infrastructure events remain included rather than "
        "being regenerated or removed. Data and code availability: the frozen benchmark, chunk metadata, prompt contract, parser, "
        "retrieval and generation outputs, analysis scripts, and SHA-256 freeze manifests are released at "
        "https://github.com/REPLACE-WITH-REPOSITORY to allow independent reproduction. The source document is publicly available [13]."
    )

    # ==========================
    # X. CONCLUSION
    # ==========================
    add_h1("X. Conclusion")
    add_p(
        "In this frozen local RAG experiment on one UIT regulation document, gold-page evidence appeared in the first five results for "
        "87.22% of answerable held-out variants. The generator made the correct structured ANSWER/ABSTAIN decision for 73.42% of records; "
        "this is a decision-level measure, not answer-correctness accuracy. In addition, 95.50% of records satisfied the structural output "
        "contract, which does not establish factual correctness or groundedness. Retrieval differences across EN, MY, and MIX were not "
        "significant, but MY had a 45.00% false-abstention rate and lower decision accuracy, including 13 false abstentions with gold "
        "evidence ranked first. This finding should be read as specific to the evaluated configuration and corpus, not as a general property "
        "of Burmese queries."
    )
    add_p(
        "These findings would have been obscured by a single aggregate score. Future work should add independent semantic annotation, "
        "starting with the 34 gold-hit false abstentions. It should also run context-swap and Burmese-abstention-string experiments, "
        "index a broader set of official documents, include more unanswerable intents, and test whether the MY pattern recurs with other "
        "models and institutions."
    )

    # ==========================
    # REFERENCES
    # ==========================
    add_h1("References")
    references = [
        "[1] P. Lewis et al., “Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks,” in Advances in Neural Information Processing Systems, vol. 33, 2020, pp. 9459–9474.",
        "[2] L. Wang, N. Yang, X. Huang, L. Yang, R. Majumder, and F. Wei, “Multilingual E5 Text Embeddings: A Technical Report,” arXiv:2402.05672, 2024, doi: 10.48550/arXiv.2402.05672.",
        "[3] S. Es, J. James, L. Espinosa Anke, and S. Schockaert, “RAGAs: Automated Evaluation of Retrieval Augmented Generation,” in Proc. 18th Conf. European Chapter of the Association for Computational Linguistics: System Demonstrations, 2024, pp. 150–158, doi: 10.18653/v1/2024.eacl-demo.16.",
        "[4] J. Do, J. Lee, and S.-w. Hwang, “ContrastiveMix: Overcoming Code-Mixing Dilemma in Cross-Lingual Transfer for Information Retrieval,” in Proc. 2024 Conf. North American Chapter of the Association for Computational Linguistics: Human Language Technologies, Volume 2 (Short Papers), 2024, pp. 197–204, doi: 10.18653/v1/2024.naacl-short.17.",
        "[5] J. Park and H. Lee, “Investigating Language Preference of Multilingual RAG Systems,” in Findings of the Association for Computational Linguistics: ACL 2025, 2025, pp. 5647–5675, doi: 10.18653/v1/2025.findings-acl.295.",
        "[6] B. Wen, J. Yao, S. Feng, C. Xu, Y. Tsvetkov, B. Howe, and L. L. Wang, “Know Your Limits: A Survey of Abstention in Large Language Models,” Transactions of the Association for Computational Linguistics, vol. 13, pp. 529–556, 2025, doi: 10.1162/tacl_a_00754.",
        "[7] L. Huang et al., “Alleviating Hallucinations from Knowledge Misalignment in Large Language Models via Selective Abstention Learning,” in Proc. 63rd Annual Meeting of the Association for Computational Linguistics, Volume 1 (Long Papers), 2025, pp. 24564–24579, doi: 10.18653/v1/2025.acl-long.1199.",
        "[8] T. A. Chang, K. Tomanek, J. Hoffmann, N. Thain, E. MacMurray van Liemt, K. Meier-Hellstern, and L. Dixon, “Detecting Hallucination and Coverage Errors in Retrieval Augmented Generation for Controversial Topics,” in Proc. 2024 Joint International Conference on Computational Linguistics, Language Resources and Evaluation, 2024, pp. 4729–4743.",
        "[9] T. Aung, J. R. Montalan, J. G. Ngui, and P. Limkonchotiwat, “BURMESE-SAN: Burmese NLP Benchmark for Evaluating Large Language Models,” in Proc. 15th Language Resources and Evaluation Conference, 2026, pp. 224–245, doi: 10.63317/54dxgzy8h77c.",
        "[10] E. Nie, S. Liang, H. Schmid, and H. Schütze, “Cross-Lingual Retrieval Augmented Prompt for Low-Resource Languages,” in Findings of the Association for Computational Linguistics: ACL 2023, 2023, pp. 8320–8340, doi: 10.18653/v1/2023.findings-acl.528.",
        "[11] A. S. Ipa, M. A. T. Rony, and M. S. Islam, “Empowering Low-Resource Languages: TraSe Architecture for Enhanced Retrieval-Augmented Generation in Bangla,” in Proc. 1st Workshop on Language Models for Underserved Communities, 2025, pp. 8–15, doi: 10.18653/v1/2025.lm4uc-1.2.",
        "[12] H. Sun, Y. Wang, and S. Zhang, “Retrieval-Augmented Generation for Domain-Specific Question Answering: A Case Study on Pittsburgh and CMU,” arXiv:2411.13691, 2024, doi: 10.48550/arXiv.2411.13691.",
        "[13] University of Information Technology, “Undergraduate Academic Regulations (4 Years Program),” June 2025, 24 pp., source ID UIT-SRC-025. [Online]. Available: https://uit.edu.mm/academic-credits/"
    ]

    for ref in references:
        p_ref = add_p(ref, style="references", space_before=1, space_after=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        p_ref.paragraph_format.left_indent = Inches(0.2)
        p_ref.paragraph_format.first_line_indent = Inches(-0.2)
        if p_ref.runs:
            p_ref.runs[0].font.size = Pt(8)

    doc.save(output_path)
    print(f"Successfully generated {output_path}!")

if __name__ == "__main__":
    build_docx()
