import docx
from docx.oxml import parse_xml

doc = docx.Document('docs/ICAIT2026-Template.docx')
b = doc._body._element

# Clear all children except body's sectPr
body_sectPr = b.xpath('./w:sectPr')[0]
for child in list(b):
    if child is not body_sectPr:
        b.remove(child)

# Title
p_title = doc.add_paragraph('Language-Conditioned Abstention in Burmese-English RAG', style='paper title')

# Table for authors (2 authors)
table_authors = doc.add_table(rows=1, cols=2)
table_authors.autofit = False

# Author 1
c0 = table_authors.rows[0].cells[0]
p0 = c0.paragraphs[0]
p0.style = 'Author'
p0.paragraph_format.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
p0.add_run('Han Nyi Min Htut\n').bold = True
p0.add_run('University of Information Technology\nYangon, Myanmar\nhannyiminhtut@uit.edu.mm')

# Author 2
c1 = table_authors.rows[0].cells[1]
p1 = c1.paragraphs[0]
p1.style = 'Author'
p1.paragraph_format.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
p1.add_run('Thet Thet Zin\n').bold = True
p1.add_run('University of Information Technology\nYangon, Myanmar\nthetthetzin@uit.edu.mm')

# Remove borders from authors table
tblPr = table_authors._tbl.tblPr
tblBorders = parse_xml(
    '<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/>'
    '<w:insideH w:val="none"/><w:insideV w:val="none"/>'
    '</w:tblBorders>'
)
tblPr.append(tblBorders)

# Section break paragraph to transition to 2 columns
p_break = doc.add_paragraph()
sect_xml = (
    '<w:sectPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:type w:val="continuous"/>'
    '<w:cols w:space="720"/>'
    '</w:sectPr>'
)
p_break._p.get_or_add_pPr().append(parse_xml(sect_xml))

# Abstract paragraph (in 2-column section)
p_abs = doc.add_paragraph(style='Abstract')
p_abs.add_run('Abstract—').bold = True
p_abs.add_run('This is a test of abstract in two columns.')

# Body text
p_body = doc.add_paragraph('This is body text in column format.', style='Body Text')

# Ensure the body_sectPr has 2 columns
body_cols = body_sectPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cols')
if body_cols is not None:
    body_cols.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num', '2')
    body_cols.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}space', '360')
else:
    body_sectPr.append(parse_xml('<w:cols xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:num="2" w:space="360"/>'))

doc.save('docs/test_2col.docx')
print('test_2col.docx created successfully!')
